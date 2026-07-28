"""PDF, DOCX, PPTX parsers for mzero."""

import os
from pathlib import Path
from mzero.types import Document, DocumentType
from mzero.utils.logger import logger


class DocParser:
    @staticmethod
    def parse_pdf(file_path: str) -> Document:
        content = ""
        metadata = {}
        try:
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            metadata["page_count"] = len(reader.pages)
            pages_text = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    pages_text.append(f"--- Page {i+1} ---\n{text}")
            content = "\n\n".join(pages_text)
        except Exception as e:
            logger.warning(f"Error parsing PDF with pypdf: {e}. Reading raw content.")
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()

        return Document(
            id=file_path,
            source_path=file_path,
            content=content,
            doc_type=DocumentType.PDF,
            metadata=metadata
        )

    @staticmethod
    def parse_docx(file_path: str) -> Document:
        content = ""
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            content = "\n\n".join(paragraphs)
        except Exception as e:
            logger.warning(f"Error parsing DOCX with python-docx: {e}")
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()

        return Document(
            id=file_path,
            source_path=file_path,
            content=content,
            doc_type=DocumentType.DOCX
        )

    @staticmethod
    def parse_pptx(file_path: str) -> Document:
        content = ""
        try:
            import pptx
            prs = pptx.Presentation(file_path)
            slides_text = []
            for i, slide in enumerate(prs.slides):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_texts.append(shape.text.strip())
                if slide_texts:
                    slides_text.append(f"--- Slide {i+1} ---\n" + "\n".join(slide_texts))
            content = "\n\n".join(slides_text)
        except Exception as e:
            logger.warning(f"Error parsing PPTX with python-pptx: {e}")
            content = f"[Presentation file: {os.path.basename(file_path)}]"

        return Document(
            id=file_path,
            source_path=file_path,
            content=content,
            doc_type=DocumentType.PPTX
        )
